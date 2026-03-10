#!/usr/bin/env python3
"""Modular budget calculation engine - can be used for any project.

Interactive mode: Prompts for project details
Config mode: Loads from JSON file
"""
from __future__ import annotations
import argparse
import json
import sys
from pathlib import Path
from typing import Dict, List, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


def format_euro(x: float) -> str:
    return f"€ {x:,.2f}"


def load_rates_from_config(config_path: Optional[Path] = None) -> Dict[str, float]:
    """Load hourly rates from config or use defaults."""
    if config_path and config_path.exists():
        data = json.loads(config_path.read_text())
        return data.get("rates", {})
    
    # Default DNV 2026 rates
    return {
        "professional": 145,
        "senior": 189,
        "principal": 231,
    }


def prompt_for_rates() -> Dict[str, float]:
    """Interactively prompt for hourly rates."""
    print("\n=== HOURLY RATES ===")
    print("Enter hourly rates in EUR (or press Enter for DNV 2026 defaults):")
    
    defaults = {"professional": 145, "senior": 189, "principal": 231}
    rates = {}
    
    for role, default in defaults.items():
        response = input(f"{role.capitalize()} rate (default €{default}): ").strip()
        rates[role] = float(response) if response else default
    
    return rates


def prompt_for_project() -> Dict:
    """Interactively build a project structure."""
    print("\n" + "="*60)
    print("NEW PROJECT SETUP")
    print("="*60)
    
    project_name = input("\nProject name: ").strip()
    
    budget_low = float(input("Budget range - LOW (€): ").strip())
    budget_high = float(input("Budget range - HIGH (€): ").strip())
    
    print("\n--- WORK PACKAGES ---")
    print("You'll now define work packages and tasks.")
    print("Each task needs hours for: Professional, Senior, Principal\n")
    
    sections = []
    section_num = 1
    
    while True:
        print(f"\n=== Work Package {section_num} ===")
        section_name = input("Work package name (or 'done' to finish): ").strip()
        
        if section_name.lower() == 'done':
            break
        
        tasks = []
        task_num = 1
        
        while True:
            print(f"\n  Task {task_num} for '{section_name}'")
            task_name = input("  Task name (or 'done' for next package): ").strip()
            
            if task_name.lower() == 'done':
                break
            
            print("  Hours by level:")
            prof_h = float(input("    Professional: ").strip() or "0")
            sr_h = float(input("    Senior: ").strip() or "0")
            pr_h = float(input("    Principal: ").strip() or "0")
            
            tasks.append({
                "name": task_name,
                "hours": {
                    "professional": prof_h,
                    "senior": sr_h,
                    "principal": pr_h
                }
            })
            task_num += 1
        
        if tasks:
            sections.append({
                "section": f"{section_num}.{section_num} {section_name}",
                "subtasks": tasks
            })
            section_num += 1
    
    return {
        "name": project_name,
        "budget_range": (budget_low, budget_high),
        "tasks": sections
    }


def load_project_from_json(json_path: Path) -> Dict:
    """Load project structure from JSON config file."""
    data = json.loads(json_path.read_text())
    return data


def save_project_template(output_path: Path):
    """Save a template JSON file for easy project creation."""
    template = {
        "name": "Example Project",
        "budget_range": [10000, 15000],
        "tasks": [
            {
                "section": "1.1 Project Management & Stakeholder Engagement",
                "subtasks": [
                    {
                        "name": "Project Management & Meetings",
                        "hours": {"professional": 12, "senior": 5, "principal": 2}
                    },
                    {
                        "name": "Stakeholder Consultation & Data Gathering",
                        "hours": {"professional": 8, "senior": 4, "principal": 0}
                    }
                ]
            },
            {
                "section": "1.2 Technical Analysis",
                "subtasks": [
                    {
                        "name": "Data Analysis & Modeling",
                        "hours": {"professional": 30, "senior": 15, "principal": 5}
                    },
                    {
                        "name": "Calculations & Simulations",
                        "hours": {"professional": 20, "senior": 10, "principal": 3}
                    }
                ]
            },
            {
                "section": "1.3 Reporting & Deliverables",
                "subtasks": [
                    {
                        "name": "Report Writing & Review",
                        "hours": {"professional": 15, "senior": 8, "principal": 2}
                    },
                    {
                        "name": "Contingency",
                        "hours": {"professional": 5, "senior": 3, "principal": 1}
                    }
                ]
            }
        ]
    }
    
    output_path.write_text(json.dumps(template, indent=2))
    print(f"\n✓ Template saved to: {output_path}")
    print("  Edit this file with your project details and use --config to load it")


def compute_task_costs(hours: Dict[str, float], rates: Dict[str, float]) -> Dict:
    """Calculate cost for a task given hours by consultant level."""
    prof_cost = hours.get("professional", 0) * rates["professional"]
    sr_cost = hours.get("senior", 0) * rates["senior"]
    pr_cost = hours.get("principal", 0) * rates["principal"]
    total_cost = prof_cost + sr_cost + pr_cost
    total_hours = hours.get("professional", 0) + hours.get("senior", 0) + hours.get("principal", 0)
    
    return {
        "hours": hours,
        "costs": {"professional": prof_cost, "senior": sr_cost, "principal": pr_cost},
        "total_cost": total_cost,
        "total_hours": total_hours
    }


def compute_project_breakdown(project: Dict, rates: Dict[str, float]) -> Dict:
    """Compute full breakdown for a project."""
    section_results = []
    total_prof_hours = 0
    total_sr_hours = 0
    total_pr_hours = 0
    total_cost = 0
    
    for section in project["tasks"]:
        section_tasks = []
        section_prof_hours = 0
        section_sr_hours = 0
        section_pr_hours = 0
        section_cost = 0
        
        for subtask in section["subtasks"]:
            task_result = compute_task_costs(subtask["hours"], rates)
            task_result["name"] = subtask["name"]
            section_tasks.append(task_result)
            
            section_prof_hours += subtask["hours"].get("professional", 0)
            section_sr_hours += subtask["hours"].get("senior", 0)
            section_pr_hours += subtask["hours"].get("principal", 0)
            section_cost += task_result["total_cost"]
        
        section_results.append({
            "section": section["section"],
            "tasks": section_tasks,
            "totals": {
                "professional": section_prof_hours,
                "senior": section_sr_hours,
                "principal": section_pr_hours,
                "cost": section_cost,
                "hours": section_prof_hours + section_sr_hours + section_pr_hours
            }
        })
        
        total_prof_hours += section_prof_hours
        total_sr_hours += section_sr_hours
        total_pr_hours += section_pr_hours
        total_cost += section_cost
    
    return {
        "project_name": project["name"],
        "budget_range": tuple(project["budget_range"]),
        "sections": section_results,
        "totals": {
            "professional": total_prof_hours,
            "senior": total_sr_hours,
            "principal": total_pr_hours,
            "total_hours": total_prof_hours + total_sr_hours + total_pr_hours,
            "total_cost": total_cost
        },
        "rates": rates
    }


def print_project_summary(result: Dict):
    """Print summary to console."""
    print(f"\n{'='*80}")
    print(f"Project: {result['project_name']}")
    print(f"Budget Range: {format_euro(result['budget_range'][0])} - {format_euro(result['budget_range'][1])}")
    print(f"Rates: Professional €{result['rates']['professional']}/h | Senior €{result['rates']['senior']}/h | Principal €{result['rates']['principal']}/h")
    print(f"{'='*80}\n")
    
    for section in result["sections"]:
        print(f"{section['section']}")
        for task in section["tasks"]:
            print(f"  {task['name']:<50} {format_euro(task['total_cost']):>12}")
        print(f"  {'Subtotal':<50} {format_euro(section['totals']['cost']):>12}")
        print()
    
    print(f"{'TOTAL':<52} {format_euro(result['totals']['total_cost']):>12}")
    print(f"Total Hours: {result['totals']['total_hours']:.1f}h\n")


def export_to_excel(result: Dict, output_path: str):
    """Export project breakdown to Excel with formulas."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget Breakdown"
    
    rates = result['rates']
    
    # Styling
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    section_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    section_font = Font(bold=True, size=10)
    subtotal_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    subtotal_font = Font(bold=True, size=10)
    total_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
    total_font = Font(bold=True, size=11)
    rate_fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
    rate_font = Font(bold=True, size=10)
    currency_format = '€ #,##0.00'
    number_format = '#,##0'
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Title rows
    ws.merge_cells('A1:E1')
    ws['A1'] = result['project_name']
    ws['A1'].font = Font(bold=True, size=14)
    
    ws.merge_cells('A2:E2')
    ws['A2'] = f"Budget Range: {format_euro(result['budget_range'][0])} - {format_euro(result['budget_range'][1])}"
    ws['A2'].font = Font(size=10)
    
    ws.merge_cells('A3:E3')
    ws['A3'] = f"Rates: Professional €{rates['professional']}/h | Senior €{rates['senior']}/h | Principal €{rates['principal']}/h"
    ws['A3'].font = Font(size=10)
    
    # Headers
    row = 5
    headers = ['Work Package / Task', 'Professional Hours', 'Senior Hours', 'Principal Hours', 'Contract Price']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
    
    # Rate row
    row += 1
    ws.cell(row=row, column=1, value="Hourly Rates →").fill = rate_fill
    ws.cell(row=row, column=1).font = rate_font
    ws.cell(row=row, column=1).border = border
    
    rate_cells = {}
    for col, (role, rate) in enumerate([('professional', rates['professional']), 
                                         ('senior', rates['senior']), 
                                         ('principal', rates['principal'])], 2):
        cell = ws.cell(row=row, column=col, value=rate)
        cell.fill = rate_fill
        cell.font = rate_font
        cell.number_format = '€ #,##0'
        cell.alignment = Alignment(horizontal='center')
        cell.border = border
        rate_cells[role] = f"{get_column_letter(col)}{row}"
    
    ws.cell(row=row, column=5, value="(Formulas calculate automatically)").fill = rate_fill
    ws.cell(row=row, column=5).font = Font(italic=True, size=9)
    ws.cell(row=row, column=5).border = border
    
    row += 1
    
    # Data rows
    section_row_map = {}
    
    for section in result["sections"]:
        # Section header
        ws.merge_cells(f'A{row}:E{row}')
        cell = ws.cell(row=row, column=1, value=section['section'])
        cell.fill = section_fill
        cell.font = section_font
        cell.border = border
        row += 1
        
        section_start_row = row
        
        # Tasks
        for task in section["tasks"]:
            ws.cell(row=row, column=1, value=f"  {task['name']}").border = border
            ws.cell(row=row, column=2, value=task["hours"].get("professional", 0)).border = border
            ws.cell(row=row, column=2).number_format = number_format
            ws.cell(row=row, column=3, value=task["hours"].get("senior", 0)).border = border
            ws.cell(row=row, column=3).number_format = number_format
            ws.cell(row=row, column=4, value=task["hours"].get("principal", 0)).border = border
            ws.cell(row=row, column=4).number_format = number_format
            
            # Formula
            formula = f"=B{row}*${rate_cells['professional']}+C{row}*${rate_cells['senior']}+D{row}*${rate_cells['principal']}"
            price_cell = ws.cell(row=row, column=5, value=formula)
            price_cell.number_format = currency_format
            price_cell.border = border
            
            row += 1
        
        section_end_row = row - 1
        
        # Subtotal
        for col in range(1, 6):
            ws.cell(row=row, column=col).fill = subtotal_fill
            ws.cell(row=row, column=col).font = subtotal_font
            ws.cell(row=row, column=col).border = border
        
        ws.cell(row=row, column=1, value="  Subtotal")
        section_row_map[section['section']] = row
        
        for col, col_letter in enumerate(['B', 'C', 'D', 'E'], 2):
            formula = f"=SUM({col_letter}{section_start_row}:{col_letter}{section_end_row})"
            cell = ws.cell(row=row, column=col, value=formula)
            if col <= 4:
                cell.number_format = number_format
            else:
                cell.number_format = currency_format
        
        row += 1
        row += 1
    
    # Grand total
    total_row = row
    for col in range(1, 6):
        ws.cell(row=row, column=col).fill = total_fill
        ws.cell(row=row, column=col).font = total_font
        ws.cell(row=row, column=col).border = border
    
    ws.cell(row=row, column=1, value="TOTAL HOURS & COST")
    
    # Grand total formulas
    subtotal_rows = list(section_row_map.values())
    for col, col_letter in enumerate(['B', 'C', 'D'], 2):
        formula = "=" + "+".join([f"{col_letter}{r}" for r in subtotal_rows])
        cell = ws.cell(row=row, column=col, value=formula)
        cell.number_format = number_format
    
    total_cost_formula = f"=B{total_row}*${rate_cells['professional']}+C{total_row}*${rate_cells['senior']}+D{total_row}*${rate_cells['principal']}"
    ws.cell(row=row, column=5, value=total_cost_formula).number_format = currency_format
    
    row += 1
    
    # Total hours
    total_hours_formula = f"=B{total_row}+C{total_row}+D{total_row}"
    ws.cell(row=row, column=1, value="Total Hours:")
    ws.cell(row=row, column=1).font = Font(bold=True)
    ws.cell(row=row, column=2, value=total_hours_formula)
    ws.cell(row=row, column=2).font = Font(bold=True)
    ws.cell(row=row, column=2).number_format = '#,##0.0'
    
    # Client summary section
    row += 3
    ws.merge_cells(f'A{row}:D{row}')
    ws.cell(row=row, column=1, value="CLIENT SUMMARY - Work Package Overview")
    ws.cell(row=row, column=1).font = Font(bold=True, size=12, color="1F4E78")
    ws.cell(row=row, column=1).fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    row += 1
    
    # Summary headers
    summary_headers = ['Work Package', 'Total Hours', 'Cost', 'Notes']
    for col, header in enumerate(summary_headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.alignment = Alignment(horizontal='center')
        cell.border = border
    row += 1
    
    summary_start_row = row
    
    # Summary rows
    for section in result["sections"]:
        subtotal_row = section_row_map.get(section["section"])
        
        ws.cell(row=row, column=1, value=section["section"].split(' ', 1)[1] if ' ' in section["section"] else section["section"])
        ws.cell(row=row, column=1).border = border
        
        if subtotal_row:
            ws.cell(row=row, column=2, value=f"=B{subtotal_row}+C{subtotal_row}+D{subtotal_row}")
            ws.cell(row=row, column=3, value=f"=E{subtotal_row}")
        
        ws.cell(row=row, column=2).number_format = '#,##0.0'
        ws.cell(row=row, column=2).border = border
        ws.cell(row=row, column=2).alignment = Alignment(horizontal='center')
        
        ws.cell(row=row, column=3).number_format = currency_format
        ws.cell(row=row, column=3).border = border
        
        # Notes
        notes = "Technical work and deliverables"
        ws.cell(row=row, column=4, value=notes)
        ws.cell(row=row, column=4).border = border
        ws.cell(row=row, column=4).font = Font(size=9, italic=True)
        
        row += 1
    
    summary_end_row = row - 1
    
    # Summary total
    for col in range(1, 5):
        ws.cell(row=row, column=col).fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
        ws.cell(row=row, column=col).font = Font(bold=True, size=11)
        ws.cell(row=row, column=col).border = border
    
    ws.cell(row=row, column=1, value="PROJECT TOTAL")
    ws.cell(row=row, column=2, value=f"=SUM(B{summary_start_row}:B{summary_end_row})")
    ws.cell(row=row, column=2).number_format = '#,##0.0'
    ws.cell(row=row, column=3, value=f"=SUM(C{summary_start_row}:C{summary_end_row})")
    ws.cell(row=row, column=3).number_format = currency_format
    
    # Column widths
    ws.column_dimensions['A'].width = 60
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 45
    ws.column_dimensions['E'].width = 18
    
    wb.save(output_path)
    print(f"\n✓ Excel file saved: {output_path}")
    print(f"  → All costs use formulas - edit hours and costs update automatically!")


def main():
    parser = argparse.ArgumentParser(description="Modular budget calculation engine")
    parser.add_argument("--config", "-c", help="Load project from JSON config file")
    parser.add_argument("--output", "-o", help="Export to Excel file")
    parser.add_argument("--template", "-t", help="Generate template JSON file")
    parser.add_argument("--interactive", "-i", action="store_true", help="Interactive mode (prompts for project details)")
    
    args = parser.parse_args()
    
    # Generate template
    if args.template:
        save_project_template(Path(args.template))
        return
    
    # Load or prompt for rates
    rates = prompt_for_rates() if args.interactive else load_rates_from_config()
    
    # Load or prompt for project
    if args.config:
        print(f"\nLoading project from: {args.config}")
        project = load_project_from_json(Path(args.config))
    elif args.interactive:
        project = prompt_for_project()
    else:
        print("\nError: Use --config to load from file OR --interactive to enter details")
        print("       Use --template <file> to generate a template JSON")
        sys.exit(1)
    
    # Compute breakdown
    result = compute_project_breakdown(project, rates)
    
    # Print summary
    print_project_summary(result)
    
    # Export to Excel
    if args.output:
        export_to_excel(result, args.output)
    else:
        print("\nTip: Use --output <file.xlsx> to export to Excel")


if __name__ == '__main__':
    main()
